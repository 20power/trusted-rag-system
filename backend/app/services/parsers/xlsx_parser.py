from __future__ import annotations

import re
import tempfile
import xml.etree.ElementTree as ET
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from app.services.parsers.common import ParsedContent

PARSER_VERSION = "0.1.0"
CELL_RANGE_PATTERN = re.compile(
    r"^\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?$"
    r"|^[A-Za-z]{1,3}:[A-Za-z]{1,3}$"
)
SPREADSHEET_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def normalize_cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _sanitize_invalid_auto_filters(source: Path, destination: Path) -> bool:
    changed = False
    with ZipFile(source, "r") as input_archive, ZipFile(
        destination,
        "w",
        compression=ZIP_DEFLATED,
    ) as output_archive:
        for info in input_archive.infolist():
            payload = input_archive.read(info.filename)
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                root = ET.fromstring(payload)
                for child in list(root):
                    if child.tag != f"{{{SPREADSHEET_NAMESPACE}}}autoFilter":
                        continue
                    reference = child.attrib.get("ref", "")
                    if not CELL_RANGE_PATTERN.fullmatch(reference):
                        root.remove(child)
                        changed = True
                if changed:
                    payload = ET.tostring(
                        root,
                        encoding="utf-8",
                        xml_declaration=True,
                    )
            output_archive.writestr(info, payload)
    return changed


@contextmanager
def _workbook_source(path: Path):
    with tempfile.TemporaryDirectory(prefix="rag-xlsx-sanitize-") as temporary:
        sanitized = Path(temporary) / "workbook.xlsx"
        changed = _sanitize_invalid_auto_filters(path, sanitized)
        yield sanitized if changed else path, changed


def parse_xlsx(path: Path, *, max_cells: int) -> ParsedContent:
    sheets: list[dict[str, Any]] = []
    total_non_empty = 0
    warnings: list[str] = []

    with _workbook_source(path) as (workbook_path, sanitized):
        if sanitized:
            warnings.append("已忽略不符合OOXML单元格范围规范的自动筛选元数据")
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            for worksheet in workbook.worksheets:
                cells: list[dict[str, Any]] = []
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        total_non_empty += 1
                        if total_non_empty > max_cells:
                            raise ValueError(
                                f"工作簿非空单元格超过安全上限 {max_cells}，"
                                "请提高 PARSER_MAX_CELLS"
                            )
                        cells.append(
                            {
                                "address": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "value": normalize_cell_value(cell.value),
                                "data_type": cell.data_type,
                                "number_format": cell.number_format,
                            }
                        )
                sheets.append(
                    {
                        "name": worksheet.title,
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                        "non_empty_cells": len(cells),
                        "cells": cells,
                    }
                )
        finally:
            workbook.close()

    if total_non_empty == 0:
        warnings.append("工作簿未发现非空单元格")
    return ParsedContent(
        parser="openpyxl",
        parser_version=PARSER_VERSION,
        source_format="xlsx",
        workbook={"sheet_count": len(sheets), "non_empty_cells": total_non_empty, "sheets": sheets},
        warnings=warnings,
    )
