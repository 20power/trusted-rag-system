import re
from pathlib import Path

from docx import Document as WordDocument

from app.services.parsers.common import ParsedContent, TextBlock

PARSER_VERSION = "0.2.0"

HEADING_STYLE_PATTERN = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)
STRUCTURAL_HEADING_PATTERN = re.compile(
    r"^第[一二三四五六七八九十百千万零〇两0-9]+([编章节])"
)
ARTICLE_PATTERN = re.compile(
    r"^(第[一二三四五六七八九十百千万零〇两0-9]+条(?:之[一二三四五六七八九十0-9]+)?)"
)
STRUCTURAL_LEVELS = {"编": 1, "章": 2, "节": 3}


def _heading_level(text: str, style: str | None) -> int | None:
    if style:
        style_match = HEADING_STYLE_PATTERN.search(style)
        if style_match:
            return int(style_match.group(1))
    structural_match = STRUCTURAL_HEADING_PATTERN.match(text)
    if structural_match:
        return STRUCTURAL_LEVELS[structural_match.group(1)]
    return None


def parse_docx(path: Path) -> ParsedContent:
    document = WordDocument(path)
    blocks: list[TextBlock] = []
    section_levels: dict[int, str] = {}
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else None
        heading_level = _heading_level(text, style)
        if heading_level is not None:
            section_levels[heading_level] = text
            section_levels = {
                level: value
                for level, value in section_levels.items()
                if level <= heading_level
            }
        section_path = [
            section_levels[level]
            for level in sorted(section_levels)
        ]
        article_match = ARTICLE_PATTERN.match(text)
        blocks.append(
            TextBlock(
                locator=f"paragraph:{index}",
                text=text,
                block_type="heading" if heading_level is not None else "paragraph",
                metadata={
                    "style": style,
                    "heading_level": heading_level,
                    "section_path": section_path,
                    "article_no": article_match.group(1) if article_match else None,
                },
            )
        )

    tables: list[dict[str, object]] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for column_index, cell in enumerate(row.cells, start=1):
                cells.append(
                    {
                        "row": row_index,
                        "column": column_index,
                        "text": cell.text.strip(),
                    }
                )
            rows.append(cells)
        tables.append({"table_index": table_index, "rows": rows})

    return ParsedContent(
        parser="python-docx",
        parser_version=PARSER_VERSION,
        source_format="docx",
        text_blocks=blocks,
        tables=tables,
    )
