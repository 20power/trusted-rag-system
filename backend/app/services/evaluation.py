from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.services.answering import build_answer
from app.services.hybrid_retrieval import retrieve_chunks
from app.services.retrieval import character_ngrams, normalize_text, search_chunks

CELL_PATTERN = re.compile(r"单元格[：:]\s*([A-Z]+\d+)", re.IGNORECASE)
SHEET_PATTERN = re.compile(r"工作表[：:]\s*([^；;]+)")


@dataclass(slots=True)
class EvaluationCase:
    question_id: str
    source_type: str
    question: str
    retrieval_query: str
    source_title: str
    expected_sheet: str | None
    expected_cell: str | None
    expected_text: str | None
    options: dict[str, str] = field(default_factory=dict)
    expected_answer: str | None = None
    expected_answer_text: str | None = None


def load_evaluation_cases(
    workbook_path: Path,
    *,
    source_type: str | None = None,
    limit: int | None = None,
) -> list[EvaluationCase]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}

    cases: list[EvaluationCase] = []
    for row in rows:
        row_source_type = str(row[positions["source_type"]] or "")
        if source_type and row_source_type != source_type:
            continue
        evidence = str(row[positions["evidence"]] or "")
        cell_match = CELL_PATTERN.search(evidence)
        sheet_match = SHEET_PATTERN.search(evidence)
        question = str(row[positions["question"]])
        options = [
            str(row[positions[column]] or "")
            for column in ("option_a", "option_b", "option_c", "option_d")
        ]
        retrieval_query = question
        if row_source_type != "excel":
            retrieval_query = f"{question}\n候选表述：" + "；".join(options)
        cases.append(
            EvaluationCase(
                question_id=str(row[positions["id"]]),
                source_type=row_source_type,
                question=question,
                retrieval_query=retrieval_query,
                source_title=str(row[positions["source_title"]]),
                expected_sheet=sheet_match.group(1).strip() if sheet_match else None,
                expected_cell=cell_match.group(1).upper() if cell_match else None,
                expected_text=evidence if row_source_type != "excel" else None,
                options={
                    letter: str(row[positions[f"option_{letter.lower()}"]] or "")
                    for letter in ("A", "B", "C", "D")
                },
                expected_answer=str(row[positions["answer"]] or "").upper() or None,
                expected_answer_text=str(row[positions["answer_text"]] or "") or None,
            )
        )
        if limit and len(cases) >= limit:
            break
    workbook.close()
    return cases


def _title_matches(actual: str, expected: str) -> bool:
    actual_normalized = normalize_text(actual)
    expected_normalized = normalize_text(expected)
    return bool(
        actual_normalized
        and expected_normalized
        and (
            actual_normalized in expected_normalized
            or expected_normalized in actual_normalized
        )
    )


def _hit_title_matches(hit: Any, expected: str) -> bool:
    return _title_matches(hit.title, expected) or _title_matches(hit.source_file, expected)


def _locator_matches(hit: Any, case: EvaluationCase) -> bool:
    if not _hit_title_matches(hit, case.source_title):
        return False
    if case.expected_sheet:
        actual_sheet = str(hit.metadata.get("sheet") or "")
        if normalize_text(actual_sheet) != normalize_text(case.expected_sheet):
            return False
    if case.expected_cell:
        addresses = {str(value).upper() for value in hit.metadata.get("addresses") or []}
        if case.expected_cell not in addresses:
            return False
    return True


def _text_evidence_rank(hits: list[Any], case: EvaluationCase) -> int | None:
    if not case.expected_text:
        return None
    facts = [
        normalize_text(value)
        for value in re.split(r"[；;]", case.expected_text)
        if normalize_text(value)
    ]
    if not facts:
        return None
    matched_ranks: list[int] = []
    for fact in facts:
        fact_grams = character_ngrams(fact)
        rank = next(
            (
                rank
                for rank, hit in enumerate(hits, start=1)
                if _hit_title_matches(hit, case.source_title)
                and (
                    fact in normalize_text(hit.content)
                    or normalize_text(hit.content) in fact
                    or (
                        fact_grams
                        and len(fact_grams & character_ngrams(hit.content))
                        / len(fact_grams)
                        >= 0.6
                    )
                )
            ),
            None,
        )
        if rank is None:
            return None
        matched_ranks.append(rank)
    return max(matched_ranks)


def evaluate_retrieval(
    db: Session,
    cases: list[EvaluationCase],
    *,
    top_k: int = 5,
    mode: str = "lexical",
    settings: Settings | None = None,
) -> dict[str, Any]:
    if mode not in {"lexical", "hybrid"}:
        raise ValueError(f"不支持的检索评测模式：{mode}")
    if mode == "hybrid" and settings is None:
        raise ValueError("混合检索评测必须提供 Settings")

    details: list[dict[str, Any]] = []
    document_hits = 0
    evidence_hits = 0
    reciprocal_ranks: list[float] = []
    actual_modes: dict[str, int] = {}

    for case in cases:
        if mode == "hybrid":
            retrieval = retrieve_chunks(
                db,
                case.retrieval_query,
                settings=settings,
                top_k=top_k,
            )
            hits = retrieval.hits
            actual_modes[retrieval.mode] = actual_modes.get(retrieval.mode, 0) + 1
        else:
            hits = search_chunks(db, case.retrieval_query, top_k=top_k)
            actual_modes["lexical"] = actual_modes.get("lexical", 0) + 1
        document_rank = next(
            (
                rank
                for rank, hit in enumerate(hits, start=1)
                if _hit_title_matches(hit, case.source_title)
            ),
            None,
        )
        if case.expected_text:
            evidence_rank = _text_evidence_rank(hits, case)
        else:
            evidence_rank = next(
                (
                    rank
                    for rank, hit in enumerate(hits, start=1)
                    if _locator_matches(hit, case)
                ),
                None,
            )
        document_hits += document_rank is not None
        evidence_hits += evidence_rank is not None
        reciprocal_ranks.append(1.0 / evidence_rank if evidence_rank else 0.0)
        details.append(
            {
                **asdict(case),
                "document_rank": document_rank,
                "evidence_rank": evidence_rank,
                "top_hit": hits[0].to_evidence() if hits else None,
            }
        )

    total = len(cases)
    return {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "requested_mode": mode,
        "actual_modes": actual_modes,
        "top_k": top_k,
        "case_count": total,
        "document_recall_at_k": document_hits / total if total else 0.0,
        "evidence_recall_at_k": evidence_hits / total if total else 0.0,
        "evidence_mrr": sum(reciprocal_ranks) / total if total else 0.0,
        "details": details,
    }


def _match_option(answer: str, options: dict[str, str]) -> str | None:
    answer_without_citations = re.sub(r"\[\d+\]", "", answer)
    explicit_letter = re.search(
        r"(?:答案|选项|应选|选择|正确项)(?:是|为)?[：:\s]*([A-D])\b",
        answer_without_citations,
        re.IGNORECASE,
    )
    if explicit_letter:
        return explicit_letter.group(1).upper()
    normalized_answer = normalize_text(answer_without_citations)
    text_matches = [
        letter
        for letter, value in options.items()
        if normalize_text(value)
        and not re.fullmatch(r"[-+]?\d+(?:\.\d+)?%?", str(value).strip())
        and normalize_text(value) in normalized_answer
    ]
    if len(text_matches) == 1:
        return text_matches[0]

    answer_numbers = {
        value.rstrip("%")
        for value in re.findall(r"[-+]?\d+(?:\.\d+)?%?", answer_without_citations)
    }
    canonical_answers = {
        format(float(number), ".15g")
        for number in answer_numbers
    }
    numeric_matches = []
    for letter, value in options.items():
        normalized_value = str(value).strip().rstrip("%")
        if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", normalized_value):
            continue
        if format(float(normalized_value), ".15g") in canonical_answers:
            numeric_matches.append(letter)
    return numeric_matches[0] if len(numeric_matches) == 1 else None


def evaluate_answers(
    db: Session,
    cases: list[EvaluationCase],
    *,
    settings: Settings,
    top_k: int = 5,
    retrieval_mode: str = "hybrid",
) -> dict[str, Any]:
    if retrieval_mode not in {"lexical", "hybrid"}:
        raise ValueError(f"不支持的检索模式：{retrieval_mode}")

    details: list[dict[str, Any]] = []
    correct = 0
    mapped = 0
    answered = 0
    citation_hits = 0
    answer_modes: Counter[str] = Counter()
    actual_retrieval_modes: Counter[str] = Counter()

    for case in cases:
        evaluation_query = case.retrieval_query
        if (
            case.source_type == "excel"
            and "哪一项" in case.question
            and "全国各地区" in case.question
        ):
            evaluation_query = (
                f"{evaluation_query}\n候选表述："
                + "；".join(case.options.values())
            )
        if retrieval_mode == "hybrid":
            retrieval = retrieve_chunks(
                db,
                evaluation_query,
                settings=settings,
                top_k=top_k,
            )
            hits = retrieval.hits
            actual_retrieval_modes[retrieval.mode] += 1
        else:
            hits = search_chunks(db, evaluation_query, top_k=top_k)
            actual_retrieval_modes["lexical"] += 1

        answer_question = case.question + "\n候选项：" + "；".join(
            f"{letter}. {value}" for letter, value in case.options.items()
        )
        result = build_answer(
            question=answer_question,
            hits=hits,
            settings=settings,
        )
        answer_modes[result.answer_mode] += 1
        answered += result.status == "answered"
        predicted = _match_option(result.answer, case.options)
        mapped += predicted is not None
        is_correct = (
            predicted is not None
            and case.expected_answer is not None
            and predicted == case.expected_answer
        )
        correct += is_correct

        cited_indices = {
            int(citation.get("evidence_index") or 0)
            for citation in result.citations
        }
        cited_hits = [
            hits[index - 1]
            for index in cited_indices
            if 1 <= index <= len(hits)
        ]
        if case.expected_text:
            citation_hit = _text_evidence_rank(cited_hits, case) is not None
        else:
            citation_hit = any(_locator_matches(hit, case) for hit in cited_hits)
        citation_hits += citation_hit

        details.append(
            {
                **asdict(case),
                "predicted_answer": predicted,
                "is_correct": is_correct,
                "status": result.status,
                "answer_mode": result.answer_mode,
                "answer": result.answer,
                "citations": result.citations,
                "warnings": result.warnings,
                "citation_hit": citation_hit,
            }
        )

    total = len(cases)
    return {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "case_count": total,
        "top_k": top_k,
        "requested_retrieval_mode": retrieval_mode,
        "actual_retrieval_modes": dict(actual_retrieval_modes),
        "answer_modes": dict(answer_modes),
        "answered_rate": answered / total if total else 0.0,
        "option_mapping_rate": mapped / total if total else 0.0,
        "accuracy": correct / total if total else 0.0,
        "citation_hit_rate": citation_hits / total if total else 0.0,
        "details": details,
    }


def write_evaluation_report(report: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".json.tmp")
    temporary.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary.replace(output_path)
