from pathlib import Path

from docx import Document

from app.services.parsers.docx_parser import parse_docx


def test_docx_parser_preserves_paragraph_and_table_locations(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("第一章 总则", level=1)
    document.add_paragraph("第一条 本办法适用于测试。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "资本充足率"
    table.cell(1, 1).text = "12%"
    document.save(source)

    parsed = parse_docx(source)

    assert [block.text for block in parsed.text_blocks] == [
        "第一章 总则",
        "第一条 本办法适用于测试。",
    ]
    assert parsed.text_blocks[1].locator == "paragraph:2"
    assert parsed.text_blocks[0].block_type == "heading"
    assert parsed.text_blocks[1].metadata["section_path"] == ["第一章 总则"]
    assert parsed.text_blocks[1].metadata["article_no"] == "第一条"
    assert parsed.tables[0]["rows"][1][1]["text"] == "12%"


def test_docx_parser_tracks_nested_heading_path_and_article_number(
    tmp_path: Path,
) -> None:
    source = tmp_path / "nested.docx"
    document = Document()
    document.add_heading("第一编 总体规则", level=1)
    document.add_heading("第二章 风险管理", level=2)
    document.add_heading("第三节 交易账簿", level=3)
    document.add_paragraph("第十二条之一 交易账簿包括短期持有的金融工具。")
    document.add_heading("第四节 银行账簿", level=3)
    document.add_paragraph("第十三条 其他头寸应划入银行账簿。")
    document.save(source)

    parsed = parse_docx(source)

    first_article = parsed.text_blocks[3]
    assert first_article.metadata["section_path"] == [
        "第一编 总体规则",
        "第二章 风险管理",
        "第三节 交易账簿",
    ]
    assert first_article.metadata["article_no"] == "第十二条之一"
    second_article = parsed.text_blocks[5]
    assert second_article.metadata["section_path"][-1] == "第四节 银行账簿"
    assert "第三节 交易账簿" not in second_article.metadata["section_path"]
